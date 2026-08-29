import io
import uuid

from sqlalchemy import select

from app.documents.defaults import build_default_document
from app.documents.numbering import compute_numbering
from app.documents.renderers.docx.renderer import render_docx
from app.documents.renderers.pdf.html_builder import build_html
from app.documents.schemas import HeadingBlock, TextRun
from app.models.group_member import GroupMember
from app.models.enums import ExportJobStatus
from app.models.export_job import ExportJob
from app.services.export_job_service import ExportJobService
from app.services.export_service import ExportService
from app.tests.conftest import OrgFixture, auth_headers
from app.tests.test_group_workflow import _internship_payload


def _setup_report_with_content(client, db, org: OrgFixture) -> tuple[str, str, str]:
    db.add(GroupMember(group_id=org.group.id, student_id=org.student.id))
    db.commit()

    teacher_token = org.teacher_token(client)
    create_resp = client.post(
        f"/api/v1/groups/{org.group.id}/internships", headers=auth_headers(teacher_token), json=_internship_payload(org.template_version.id)
    )
    internship_id = create_resp.json()["id"]
    client.post(f"/api/v1/internships/{internship_id}/publish", headers=auth_headers(teacher_token))

    student_token = org.student_token(client)
    reports = client.get("/api/v1/reports", headers=auth_headers(student_token)).json()
    report_id = reports[0]["id"]

    document_response = client.get(f"/api/v1/reports/{report_id}/document", headers=auth_headers(student_token)).json()
    doc = document_response["document"]
    section = doc["sections"][0]
    block = section["blocks"][0]
    filled = [
        {"type": "paragraph", "id": block["id"], "style_name": "Normal", "style_override": None,
         "runs": [{"kind": "text", "id": "r1", "text": "Өндірістік практика барысында студент білім алды.", "bold": False, "italic": False, "underline": False}]}
    ]
    client.patch(
        f"/api/v1/reports/{report_id}/document",
        headers=auth_headers(student_token),
        json={"expected_revision": document_response["revision"], "sections": {section["id"]: filled}},
    )

    return teacher_token, student_token, report_id


def _request_export(client, token: str, report_id: str, export_format: str = "docx"):
    return client.post(f"/api/v1/reports/{report_id}/exports/{export_format}", headers=auth_headers(token))


def _complete_export(db, job_id: str) -> ExportJob:
    service = ExportJobService(db)
    job = service.claim(uuid.UUID(job_id), "test-worker")
    assert job is not None
    data, content_type = ExportService(db).render_for_job(job)
    service.storage.save(job.storage_key, io.BytesIO(data), content_type)
    finished = service.finish(job.id, ExportJobStatus.SUCCEEDED, content_type=content_type, size_bytes=len(data))
    assert finished is not None
    return finished


class TestDocxExport:
    def test_docx_generator_omits_section_and_heading_prefixes(self):
        source = build_default_document()
        source.sections[0].blocks.append(HeadingBlock(level=2, runs=[TextRun(text="Подраздел")]))
        payload = render_docx(source, compute_numbering(source), lambda _file_id: b"").getvalue()

        import docx

        document = docx.Document(io.BytesIO(payload))
        headings = [paragraph for paragraph in document.paragraphs if paragraph.style.name.startswith("Heading")]
        assert [paragraph.text for paragraph in headings[:2]] == ["Введение", "Подраздел"]
        assert all(not paragraph._p.xpath("./w:pPr/w:numPr") for paragraph in headings[:2])

    def test_student_can_export_own_report_as_valid_docx(self, client, db, org_a: OrgFixture):
        _teacher_token, student_token, report_id = _setup_report_with_content(client, db, org_a)
        created = _request_export(client, student_token, report_id)
        assert created.status_code == 202
        job = _complete_export(db, created.json()["job_id"])
        status_response = client.get(f"/api/v1/export-jobs/{job.id}", headers=auth_headers(student_token))
        assert status_response.status_code == 200
        assert status_response.json()["status"] == "succeeded"
        resp = client.get(f"/api/v1/export-jobs/{job.id}/download", headers=auth_headers(student_token))
        assert resp.status_code == 200
        assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        assert len(resp.content) > 0

        import docx

        document = docx.Document(io.BytesIO(resp.content))
        all_text = "\n".join(p.text for p in document.paragraphs)
        assert len(document.paragraphs) > 0
        assert "Өндірістік практика барысында студент білім алды." in all_text
        assert "і" in all_text or "І" in all_text

    def test_docx_headings_have_no_generated_section_numbering(self, client, db, org_a: OrgFixture):
        _teacher_token, student_token, report_id = _setup_report_with_content(client, db, org_a)
        created = _request_export(client, student_token, report_id)
        resp = client.get(f"/api/v1/export-jobs/{_complete_export(db, created.json()['job_id']).id}/download", headers=auth_headers(student_token))
        assert resp.status_code == 200

        import docx

        document = docx.Document(io.BytesIO(resp.content))
        headings = [paragraph for paragraph in document.paragraphs if paragraph.style.name.startswith("Heading")]
        assert headings
        assert headings[0].text == "Введение"
        # A heading must not carry a Word numPr reference either. This keeps
        # Word from adding its own grey outline number when the file opens.
        assert not headings[0]._p.xpath("./w:pPr/w:numPr")

    def test_teacher_can_export_report_in_their_group(self, client, db, org_a: OrgFixture):
        teacher_token, _student_token, report_id = _setup_report_with_content(client, db, org_a)
        created = _request_export(client, teacher_token, report_id)
        assert created.status_code == 202
        job = _complete_export(db, created.json()["job_id"])
        assert client.get(f"/api/v1/export-jobs/{job.id}/download", headers=auth_headers(teacher_token)).status_code == 200

    def test_cross_tenant_cannot_export(self, client, db, org_a: OrgFixture, org_b: OrgFixture):
        _teacher_token, _student_token, report_id = _setup_report_with_content(client, db, org_a)
        org_b_teacher_token = org_b.teacher_token(client)
        resp = _request_export(client, org_b_teacher_token, report_id)
        assert resp.status_code == 404

    def test_unauthenticated_cannot_export(self, client, db, org_a: OrgFixture):
        _teacher_token, _student_token, report_id = _setup_report_with_content(client, db, org_a)
        resp = client.post(f"/api/v1/reports/{report_id}/exports/docx")
        assert resp.status_code == 401

    def test_other_students_report_not_exportable_by_different_student(self, client, db, org_a: OrgFixture):
        from app.core.security import hash_password
        from app.models.membership import OrganizationMembership
        from app.models.student import Student
        from app.models.user import User

        _teacher_token, _student_token, report_id = _setup_report_with_content(client, db, org_a)

        intruder_user = User(email="intruder@org-a.edu", hashed_password=hash_password("Practice123!"), full_name="Intruder")
        db.add(intruder_user)
        db.flush()
        intruder_membership = OrganizationMembership(user_id=intruder_user.id, organization_id=org_a.org.id, role_id=org_a.role_student.id)
        db.add(intruder_membership)
        db.flush()
        db.add(Student(membership_id=intruder_membership.id))
        db.commit()

        intruder_token = org_a.login(client, intruder_user.email)
        resp = _request_export(client, intruder_token, report_id)
        assert resp.status_code == 403


class TestPdfExport:
    def test_pdf_html_omits_generated_section_and_heading_numbers(self):
        document = build_default_document()
        document.sections[0].blocks.append(HeadingBlock(level=2, runs=[TextRun(text="Подраздел")]))
        html = build_html(document, compute_numbering(document), lambda _file_id: b"")

        assert "<h1>Введение</h1>" in html
        assert "<h2>Подраздел</h2>" in html
        assert "pf-number" not in html

    def test_export_produces_valid_pdf_with_expected_text(self, client, db, org_a: OrgFixture):
        _teacher_token, student_token, report_id = _setup_report_with_content(client, db, org_a)
        created = _request_export(client, student_token, report_id, "pdf")
        job = _complete_export(db, created.json()["job_id"])
        resp = client.get(f"/api/v1/export-jobs/{job.id}/download", headers=auth_headers(student_token))
        assert resp.status_code == 200
        assert resp.headers["content-type"] == "application/pdf"
        assert resp.content.startswith(b"%PDF")
        assert len(resp.content) > 100

    def test_cross_tenant_cannot_export_pdf(self, client, db, org_a: OrgFixture, org_b: OrgFixture):
        _teacher_token, _student_token, report_id = _setup_report_with_content(client, db, org_a)
        org_b_student_token = org_b.student_token(client)
        resp = _request_export(client, org_b_student_token, report_id, "pdf")
        assert resp.status_code == 404

    def test_export_audit_logged(self, client, db, org_a: OrgFixture):
        from app.models.audit_log import AuditLog
        from app.models.enums import AuditEventType

        _teacher_token, student_token, report_id = _setup_report_with_content(client, db, org_a)
        created = _request_export(client, student_token, report_id, "pdf")
        _complete_export(db, created.json()["job_id"])

        entries = db.execute(select(AuditLog).where(AuditLog.event_type == AuditEventType.PDF_GENERATED)).scalars().all()
        assert len(entries) == 1

    def test_active_export_is_deduplicated(self, client, db, org_a: OrgFixture):
        _teacher_token, student_token, report_id = _setup_report_with_content(client, db, org_a)
        first = _request_export(client, student_token, report_id)
        second = _request_export(client, student_token, report_id)
        assert first.status_code == second.status_code == 202
        assert first.json()["job_id"] == second.json()["job_id"]
        assert second.json()["reused_active_job"] is True

    def test_status_and_download_remain_tenant_scoped(self, client, db, org_a: OrgFixture, org_b: OrgFixture):
        _teacher_token, student_token, report_id = _setup_report_with_content(client, db, org_a)
        job = _complete_export(db, _request_export(client, student_token, report_id).json()["job_id"])
        intruder = org_b.student_token(client)
        assert client.get(f"/api/v1/export-jobs/{job.id}", headers=auth_headers(intruder)).status_code == 404
        assert client.get(f"/api/v1/export-jobs/{job.id}/download", headers=auth_headers(intruder)).status_code == 404

    def test_queued_export_can_be_cancelled(self, client, db, org_a: OrgFixture):
        _teacher_token, student_token, report_id = _setup_report_with_content(client, db, org_a)
        created = _request_export(client, student_token, report_id)
        cancelled = client.delete(f"/api/v1/export-jobs/{created.json()['job_id']}", headers=auth_headers(student_token))
        assert cancelled.status_code == 200
        assert cancelled.json()["status"] == "cancelled"
